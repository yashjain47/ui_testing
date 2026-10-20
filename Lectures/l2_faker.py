from faker import Faker
from openpyxl import Workbook

fake_data = Faker()

# wb = Workbook()
# ws = wb.active

# for i in range(1,11):

#     ws.cell(row=i,column=1).value=fake_data.name()
#     ws.cell(row=i,column=2).value=fake_data.phone_number()
#     print("Name : ",fake_data.name())
#     print("Mobile Number : ",fake_data.phone_number())


# wb.save("/home/vvdn/Desktop/Selenium/test.xlsx")

print(fake_data.currency())
