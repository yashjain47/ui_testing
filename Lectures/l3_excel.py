'''
Openpyxl

Column = A B C D ....Z

row = 1 2 3 4 5 .....

cell =  A1 B1 C1 D1 .....
        A2 B2 C2 D2 .....
        .................
        .................


'''

from openpyxl import Workbook,load_workbook




# For Saving Excel File

# wb = Workbook()
# ws = wb.active

# ws['A1'] = "Names"
# ws['B1'] = "Mobile Number"

# -------------------------------------------------------------------------------
# test_data = [["Yash","92323232323"],["Tushar","12345234523"]]

# for data in test_data:
#     ws.append(data)

# -------------------------------------------------------------------------------
# for i in range(1,6):
#     for j in range(1,5):
#         ws.cell(row=i,column=j).value = i+j

# -------------------------------------------------------------------------------




# wb.save("/home/vvdn/Desktop/Selenium/test_excel.xlsx")


########################################################################################################


# For Loading Excel File

wb = load_workbook(filename="test.xlsx")
ws = wb.active

# print(ws['A1'].value)
# print(ws.cell(row=2,column=2).value)
# print(wb['Sheet']["A1"].value)

row_ct = ws.max_row
col_ct = ws.max_column

for i in range(1,row_ct+1):
    for j in range(1,col_ct+1):
        print(ws.cell(row=i,column=j).value ,end=" ")
    print("\n")